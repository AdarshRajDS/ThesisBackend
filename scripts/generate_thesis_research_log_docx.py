#!/usr/bin/env python3
"""
Generate THESIS_RESEARCH_LOG.docx — living appendix of approaches, experiments, and design notes.
Re-run after major milestones: python scripts/generate_thesis_research_log_docx.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt


def _add_para(doc, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    return p


def build_document() -> Document:
    doc = Document()
    title = doc.add_heading("Multimodal anatomy RAG — research log & design appendix", 0)
    title.runs[0].font.size = Pt(18)

    _add_para(
        doc,
        "Auto-generated scaffold for thesis documentation. Append dated experiment entries, "
        "screenshots, and metric tables as you run studies.",
    )

    sections = [
        (
            1,
            "1. Research question & scope",
            [
                "Goal: support anatomy Q&A grounded in instructor PDFs and extracted figures, "
                "with retrievable evidence and deployable backend (FastAPI) + optional Gradio frontend.",
                "Non-goals (phase 1): formal clinical decision support; Blender MCP deferred.",
            ],
        ),
        (
            1,
            "2. Architecture (before vs after key pivots)",
            [
                "Before: single CLIP Chroma store; LLM context often empty when collection was image-heavy.",
                "After: dual retrieval — (A) MiniLM text index on data/processed/chroma from PDF ingestion; "
                "(B) CLIP multimodal index on data/processed/multimodal_chroma for images + captions.",
                "Fusion: concatenate deduplicated text passages + optional multimodal text snippets into one LLM prompt.",
            ],
        ),
        (
            1,
            "3. Text retrieval policy (current)",
            [
                "Fetch up to 24 candidates from text Chroma (similarity order preserved).",
                "Deduplicate by: (1) SHA-256 fingerprint of normalized first ~220 chars of chunk body; "
                "(2) pair (normalized PDF stem, page) where stem strips punctuation/plus so "
                "'anatomy+phys…' and 'anatomyphys…' collapse — reduces duplicate rows from double ingestion.",
                "Cap unique passages at 3 (dynamic 0..3). Only unique passages appear in `sources` and in context.",
            ],
        ),
        (
            1,
            "4. Image retrieval & URLs",
            [
                "CLIP query embedding; similarity cutoff with fallback so reranker always has candidates.",
                "Caption + CLIP cosine blend when local file exists; caption-only when remote-only (object_key).",
                "Presigned Supabase / MinIO URLs; local copy fallback to /outputs.",
                "Image dedupe: normalize object_key (and local filename) to alphanumeric-only key so logically "
                "same asset under two key variants does not appear twice.",
            ],
        ),
        (
            1,
            "5. Grounding & provenance (API metadata, not oracle)",
            [
                "The LLM almost always paraphrases unless forced to quote; distinguishing "
                "\"direct extract\" vs \"synthesized from context\" vs \"pure world knowledge\" "
                "requires explicit prompting, span-level citation, or post-hoc entailment checks — not fully solved here.",
                "`grounding.primary_basis`: retrieved_corpus_synthesized vs general_knowledge_fallback "
                "(based on whether any retrieved text/multimodal excerpt reached the prompt).",
                "`provenance_explanation`: human-readable clarification of synthesis vs fallback.",
                "`evaluation_note`: pointers to human rubric, RAGAS-style metrics, gold Q&A set, cautious LLM-as-judge.",
            ],
        ),
        (
            1,
            "6. Correctness & evaluation experiments (thesis-sized checklist)",
            [
                "Empty-corpus vs ingested-corpus A/B: measure answer length, citation overlap with gold spans, nDCG@k on retrieval.",
                "Faithfulness: for each claim sentence, binary \"supported by union of retrieved chunks\" (student annotators or lightweight NLI).",
                "Hallucination rate on held-out questions with known unsupported traps.",
                "Latency & cost: tokens in/out, embedding batch sizes, GPU vs CPU.",
                "Ablation: text-only vs multimodal-only vs fused; k and dedupe caps; rerank weights 0.25/0.75.",
                "Duplicate PDF hygiene: dedupe in retrieval does not remove duplicate vectors in DB — recommend "
                "ingestion-time file hash or canonical filename policy.",
            ],
        ),
        (
            1,
            "7. Object storage & deployment",
            [
                "STORAGE_PROVIDER minio|supabase; auto-supabase when URL + service key set.",
                "put_object uses raw bytes for Supabase; presigned GET for frontend.",
                "CORS ALLOWED_ORIGINS for HF Space + localhost.",
            ],
        ),
        (
            1,
            "8. Experiment log (template — copy row per run)",
            [
                "Date | Hypothesis | Config (k, dedupe, model) | Dataset | Metric | Result | Notes",
                "________________________________________________________________________________",
                "|                  |                     |         |      |        |      |",
            ],
        ),
        (
            1,
            "9. Known limitations",
            [
                "Generative models confabulate; retrieval does not guarantee factual answers.",
                "CLIP text vs MiniLM live in different embedding spaces — fusion is late (concat context), not joint ranking.",
                "OCR for bitmap-only figures not in text pipeline unless added explicitly.",
            ],
        ),
        (
            1,
            "10. Future work",
            [
                "Structured answers with mandatory citations [source:page] per bullet.",
                "RAGAS or TruLens in CI on a frozen eval set.",
                "OCR + figure captions into text index.",
                "Local LLM option for offline demos.",
                "Blender MCP for 3D assets when scope allows.",
            ],
        ),
    ]

    for level, heading, paras in sections:
        doc.add_heading(heading, level)
        for line in paras:
            _add_para(doc, line)

    doc.add_page_break()
    doc.add_heading("Appendix A — Raw implementation notes", 1)
    _add_para(
        doc,
        "Constants (code): _CHROMA_TEXT_CANDIDATES=24, _MAX_UNIQUE_TEXT_PASSAGES=3; "
        "image cap MAX_IMAGES=3; rerank pool then top image_metas[:5] before URL build.",
    )
    _add_para(
        doc,
        "Regenerate this file: python scripts/generate_thesis_research_log_docx.py",
        bold=True,
    )

    return doc


def main():
    root = Path(__file__).resolve().parents[1]
    out = root / "THESIS_RESEARCH_LOG.docx"
    doc = build_document()
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
